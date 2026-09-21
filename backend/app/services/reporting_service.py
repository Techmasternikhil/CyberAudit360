import os
import io
from typing import List
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
from app.models.finding import Finding
from app.models.audit import Audit
from app.services.risk_engine import RiskEngine
from app.services.framework_mapper import FrameworkMapper

def set_cell_background(cell, hex_color: str):
    """Sets the background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets padding for a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

class ReportingService:
    @staticmethod
    def get_report_text(audit: Audit, findings: List[Finding]) -> str:
        if audit is None:
            audit = Audit(
                id="live-assessment",
                name="CyberAudit360 Live Assessment",
                organization="Local Infrastructure",
                auditor="CyberAudit360 Automated Engine"
            )
        
        security_score_data = RiskEngine.calculate_overall_security_score(findings)
        coverage = FrameworkMapper.calculate_coverage(findings)
        
        report_content = f"""# Cybersecurity Audit Report: {audit.name}

## 1. Executive Summary
**Organization:** {audit.organization}
**Date:** {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S UTC')}
**Auditor:** {audit.auditor}
**Overall Security Score:** {security_score_data['score']}/100

### Findings Summary
- **Critical:** {sum(1 for f in findings if f.severity == 'CRITICAL')}
- **High:** {sum(1 for f in findings if f.severity == 'HIGH')}
- **Medium:** {sum(1 for f in findings if f.severity == 'MEDIUM')}
- **Low:** {sum(1 for f in findings if f.severity == 'LOW')}

### Framework Coverage (Estimated)
- **NIST CSF 2.0:** {coverage.get('NIST CSF 2.0', 0)}%
- **CIS Controls v8.1:** {coverage.get('CIS Controls v8.1', 0)}%
- **ISO/IEC 27001:** {coverage.get('ISO 27001', 0)}%

---
## 2. Detailed Findings
"""
        for f in findings:
            mappings = FrameworkMapper.map_finding(f.title, f.description)
            mapping_str = ", ".join([f"{k}: {v}" for k, v in mappings.items()])
            sev_val = f.severity.value if hasattr(f.severity, "value") else str(f.severity)
            stat_val = f.status.value if hasattr(f.status, "value") else str(f.status)
            
            report_content += f"""
### {f.id} - {f.title}
**Severity:** {sev_val} | **Status:** {stat_val} | **Risk Score:** {f.risk_score or 'N/A'}
**Description:**
{f.description}

**Recommendation:**
{f.recommendation or 'Verify port exposure and apply firewall restrictions.'}

**Framework Mapping:**
{mapping_str}

---
"""
        report_content += "\n## 3. Conclusion & Next Steps\nPlease review the highest severity findings and assign remediation owners immediately.\n"
        return report_content

    @staticmethod
    def generate_markdown_report(audit: Audit, findings: List[Finding], output_dir: str = "reports") -> str:
        """Generates a professional markdown audit report and writes to file."""
        os.makedirs(output_dir, exist_ok=True)
        report_content = ReportingService.get_report_text(audit, findings)
        audit_id = audit.id if audit else "assessment"
        file_path = os.path.join(output_dir, f"audit_report_{audit_id}.md")
        with open(file_path, "w", encoding="utf-8") as file:
            file.write(report_content)
        return file_path

    @staticmethod
    def generate_docx_stream(audit: Audit, findings: List[Finding]) -> io.BytesIO:
        """Generates an executive-grade Word (.docx) cybersecurity audit document."""
        if audit is None:
            audit = Audit(
                id="live-assessment",
                name="Comprehensive Cybersecurity Audit",
                organization="Internal Infrastructure & Workstations",
                auditor="CyberAudit360 Automated Engine",
                scope="Local Network Endpoints & Active Services"
            )
        
        doc = docx.Document()
        
        # Configure margins (1 inch)
        sections = doc.sections
        for s in sections:
            s.top_margin = Inches(1)
            s.bottom_margin = Inches(1)
            s.left_margin = Inches(1)
            s.right_margin = Inches(1)

        # Palette colors
        C_PRIMARY = RGBColor(15, 23, 42)      # Deep Navy/Slate (#0F172A)
        C_ACCENT = RGBColor(37, 99, 235)      # Indigo/Blue (#2563EB)
        C_MUTED = RGBColor(100, 116, 139)     # Slate Gray (#64748B)

        # 1. Document Title Header
        title_p = doc.add_paragraph()
        title_run = title_p.add_run("CyberAudit360")
        title_run.font.size = Pt(26)
        title_run.font.bold = True
        title_run.font.color.rgb = C_ACCENT
        title_p.paragraph_format.space_after = Pt(2)

        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run("EXECUTIVE CYBERSECURITY AUDIT & ASSURANCE REPORT")
        sub_run.font.size = Pt(13)
        sub_run.font.bold = True
        sub_run.font.color.rgb = C_PRIMARY
        sub_p.paragraph_format.space_after = Pt(18)

        # 2. Assessment Metadata Table
        meta_table = doc.add_table(rows=4, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False
        
        meta_data = [
            ("Target Organization", audit.organization or "Internal Assessment"),
            ("Lead Auditor / Team", audit.auditor or "CyberAudit360 Engine"),
            ("Assessment Scope", audit.scope or "Local Workstation & Network Services"),
            ("Date of Assessment", datetime.utcnow().strftime("%B %d, %Y (%H:%M UTC)"))
        ]

        for i, (label, val) in enumerate(meta_data):
            row = meta_table.rows[i]
            c0, c1 = row.cells[0], row.cells[1]
            c0.width = Inches(2.2)
            c1.width = Inches(4.3)
            
            p0 = c0.paragraphs[0]
            r0 = p0.add_run(label)
            r0.font.bold = True
            r0.font.size = Pt(9.5)
            r0.font.color.rgb = RGBColor(71, 85, 105)
            
            p1 = c1.paragraphs[0]
            r1 = p1.add_run(val)
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = C_PRIMARY
            
            set_cell_background(c0, "F1F5F9")
            set_cell_background(c1, "F8FAFC")
            set_cell_margins(c0, 80, 80, 120, 120)
            set_cell_margins(c1, 80, 80, 120, 120)

        doc.add_paragraph().paragraph_format.space_after = Pt(14)

        # 3. Section 1: Executive Summary & Score
        h1 = doc.add_heading(level=1)
        h1_run = h1.add_run("1. Executive Summary & Security Posture")
        h1_run.font.color.rgb = C_PRIMARY
        h1_run.font.size = Pt(16)
        h1.paragraph_format.space_before = Pt(12)
        h1.paragraph_format.space_after = Pt(8)

        sec_score_data = RiskEngine.calculate_overall_security_score(findings)
        score = sec_score_data['score']

        p_summary = doc.add_paragraph(
            "This report summarizes the defensive cybersecurity assessment conducted across authorized assets. "
            "The findings, exposures, and compliance postures below are calculated in accordance with standardized "
            "vulnerability scoring mechanisms (CVSS v3.1), asset business criticalities, and industry-recognized controls."
        )
        p_summary.paragraph_format.space_after = Pt(12)

        # Scorecard Highlight Box
        score_table = doc.add_table(rows=1, cols=2)
        score_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        c_left, c_right = score_table.rows[0].cells[0], score_table.rows[0].cells[1]
        c_left.width = Inches(3.0)
        c_right.width = Inches(3.5)

        p_sc = c_left.paragraphs[0]
        r_sc_lbl = p_sc.add_run("OVERALL SECURITY SCORE\n")
        r_sc_lbl.font.size = Pt(10)
        r_sc_lbl.font.bold = True
        r_sc_lbl.font.color.rgb = C_MUTED

        r_sc_val = p_sc.add_run(f"{score}")
        r_sc_val.font.size = Pt(36)
        r_sc_val.font.bold = True
        if score >= 80:
            r_sc_val.font.color.rgb = RGBColor(16, 185, 129) # Green
        elif score >= 60:
            r_sc_val.font.color.rgb = RGBColor(245, 158, 11) # Amber
        else:
            r_sc_val.font.color.rgb = RGBColor(239, 68, 68)  # Red

        r_sc_total = p_sc.add_run(" / 100")
        r_sc_total.font.size = Pt(14)
        r_sc_total.font.color.rgb = C_MUTED

        p_rt = c_right.paragraphs[0]
        crit_n = sum(1 for f in findings if f.severity == 'CRITICAL' and f.status != 'RESOLVED')
        high_n = sum(1 for f in findings if f.severity == 'HIGH' and f.status != 'RESOLVED')
        med_n = sum(1 for f in findings if f.severity == 'MEDIUM' and f.status != 'RESOLVED')
        low_n = sum(1 for f in findings if (f.severity in ['LOW', 'INFORMATIONAL']) and f.status != 'RESOLVED')
        
        p_rt.add_run("Active Exposure Overview:\n").font.bold = True
        p_rt.add_run(f"• {crit_n} Critical Vulnerabilities\n")
        p_rt.add_run(f"• {high_n} High Severity Findings\n")
        p_rt.add_run(f"• {med_n} Medium Severity Findings\n")
        p_rt.add_run(f"• {low_n} Low / Informational Items")
        p_rt.runs[0].font.size = Pt(9.5)
        for r in p_rt.runs[1:]:
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(51, 65, 85)

        set_cell_background(c_left, "F8FAFC")
        set_cell_background(c_right, "F1F5F9")
        set_cell_margins(c_left, 100, 100, 140, 140)
        set_cell_margins(c_right, 100, 100, 140, 140)

        doc.add_paragraph().paragraph_format.space_after = Pt(14)

        # 4. Section 2: Framework Alignment Summary
        coverage = FrameworkMapper.calculate_coverage(findings)
        h1_fw = doc.add_heading(level=1)
        h1_fw.add_run("2. Compliance & Regulatory Control Coverage").font.color.rgb = C_PRIMARY
        h1_fw.paragraph_format.space_before = Pt(12)
        h1_fw.paragraph_format.space_after = Pt(8)

        fw_table = doc.add_table(rows=4, cols=3)
        fw_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Header Row
        hdr_cells = fw_table.rows[0].cells
        hdr_cells[0].paragraphs[0].add_run("Framework Standard").font.bold = True
        hdr_cells[1].paragraphs[0].add_run("Scope / Focus").font.bold = True
        hdr_cells[2].paragraphs[0].add_run("Estimated Alignment").font.bold = True
        for c in hdr_cells:
            c.paragraphs[0].runs[0].font.size = Pt(9.5)
            c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
            set_cell_background(c, "1E293B")
            set_cell_margins(c, 90, 90, 120, 120)

        fw_data = [
            ("NIST CSF 2.0", "Identify, Protect, Detect, Respond, Recover", f"{coverage.get('NIST CSF 2.0', 0)}%"),
            ("CIS Controls v8.1", "Essential Cyber Hygiene & Asset Defense", f"{coverage.get('CIS Controls v8.1', 0)}%"),
            ("ISO/IEC 27001", "Information Security Management System (ISMS)", f"{coverage.get('ISO 27001', 0)}%")
        ]

        for idx, (fw_name, fw_desc, fw_val) in enumerate(fw_data):
            row = fw_table.rows[idx + 1]
            row.cells[0].paragraphs[0].add_run(fw_name).font.bold = True
            row.cells[1].paragraphs[0].add_run(fw_desc)
            r_val = row.cells[2].paragraphs[0].add_run(fw_val)
            r_val.font.bold = True
            r_val.font.color.rgb = C_ACCENT
            
            for c in row.cells:
                c.paragraphs[0].runs[0].font.size = Pt(9)
                set_cell_background(c, "FFFFFF" if idx % 2 == 0 else "F8FAFC")
                set_cell_margins(c, 80, 80, 120, 120)

        doc.add_paragraph().paragraph_format.space_after = Pt(14)

        # 5. Section 3: Detailed Findings
        h1_findings = doc.add_heading(level=1)
        h1_findings.add_run("3. Detailed Vulnerability & Risk Findings").font.color.rgb = C_PRIMARY
        h1_findings.paragraph_format.space_before = Pt(14)
        h1_findings.paragraph_format.space_after = Pt(10)

        if not findings:
            doc.add_paragraph("No active findings recorded in this audit session.")
        else:
            for idx, f in enumerate(findings, 1):
                sev_val = f.severity.value if hasattr(f.severity, "value") else str(f.severity)
                stat_val = f.status.value if hasattr(f.status, "value") else str(f.status)

                # Finding Title Header
                h2 = doc.add_heading(level=2)
                h2_run = h2.add_run(f"3.{idx} [{f.id}] {f.title}")
                h2_run.font.size = Pt(12)
                h2_run.font.bold = True
                h2_run.font.color.rgb = C_PRIMARY
                h2.paragraph_format.space_before = Pt(10)
                h2.paragraph_format.space_after = Pt(4)

                # Metadata Mini-Table
                f_table = doc.add_table(rows=2, cols=4)
                f_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                f_table.autofit = False

                labels = [
                    ("Severity", sev_val),
                    ("Status", stat_val),
                    ("Risk Score", str(f.risk_score or "N/A")),
                    ("Classification", "DEMO DATA" if f.is_demo else "LIVE AUDIT")
                ]

                # Row 0: Labels
                for col_idx, (lbl, _) in enumerate(labels):
                    cell = f_table.rows[0].cells[col_idx]
                    p = cell.paragraphs[0]
                    r = p.add_run(lbl)
                    r.font.bold = True
                    r.font.size = Pt(8.5)
                    r.font.color.rgb = RGBColor(100, 116, 139)
                    set_cell_background(cell, "F1F5F9")
                    set_cell_margins(cell, 60, 60, 80, 80)

                # Row 1: Values
                for col_idx, (_, val) in enumerate(labels):
                    cell = f_table.rows[1].cells[col_idx]
                    p = cell.paragraphs[0]
                    r = p.add_run(val)
                    r.font.bold = True
                    r.font.size = Pt(9.5)
                    
                    # Highlight severity
                    if col_idx == 0:
                        if sev_val == "CRITICAL":
                            r.font.color.rgb = RGBColor(185, 28, 28)
                        elif sev_val == "HIGH":
                            r.font.color.rgb = RGBColor(194, 65, 12)
                        elif sev_val == "MEDIUM":
                            r.font.color.rgb = RGBColor(180, 83, 9)
                        else:
                            r.font.color.rgb = RGBColor(29, 78, 216)
                    elif col_idx == 1:
                        if val == "RESOLVED":
                            r.font.color.rgb = RGBColor(16, 185, 129)
                        else:
                            r.font.color.rgb = RGBColor(225, 29, 72)
                    
                    set_cell_background(cell, "FFFFFF")
                    set_cell_margins(cell, 60, 60, 80, 80)

                # Description & Remediation paragraphs
                p_desc = doc.add_paragraph()
                p_desc.paragraph_format.space_before = Pt(6)
                p_desc.paragraph_format.space_after = Pt(3)
                lbl_desc = p_desc.add_run("Description: ")
                lbl_desc.font.bold = True
                lbl_desc.font.size = Pt(9.5)
                r_desc = p_desc.add_run(f.description or "No description recorded.")
                r_desc.font.size = Pt(9.5)

                p_rem = doc.add_paragraph()
                p_rem.paragraph_format.space_before = Pt(2)
                p_rem.paragraph_format.space_after = Pt(3)
                lbl_rem = p_rem.add_run("Recommended Action: ")
                lbl_rem.font.bold = True
                lbl_rem.font.size = Pt(9.5)
                lbl_rem.font.color.rgb = RGBColor(4, 120, 87)
                r_rem = p_rem.add_run(f.recommendation or "Inspect active listening sockets and enforce least-privilege firewall rules.")
                r_rem.font.size = Pt(9.5)

                # Compliance mapping
                mappings = FrameworkMapper.map_finding(f.title, f.description)
                p_map = doc.add_paragraph()
                p_map.paragraph_format.space_before = Pt(2)
                p_map.paragraph_format.space_after = Pt(12)
                lbl_map = p_map.add_run("Framework Mapping: ")
                lbl_map.font.bold = True
                lbl_map.font.size = Pt(9)
                lbl_map.font.color.rgb = C_MUTED
                map_str = " | ".join([f"{k}: {v}" for k, v in mappings.items()])
                r_map = p_map.add_run(map_str)
                r_map.font.size = Pt(9)
                r_map.font.italic = True

        # 6. Section 4: Sign-off & Chain-of-Custody
        h1_sign = doc.add_heading(level=1)
        h1_sign.add_run("4. Audit Authorization & Chain-of-Custody").font.color.rgb = C_PRIMARY
        h1_sign.paragraph_format.space_before = Pt(14)
        h1_sign.paragraph_format.space_after = Pt(8)

        p_audit_trail = doc.add_paragraph(
            "This document constitutes an official automated security evaluation produced by CyberAudit360. "
            "All supporting evidence artifacts have been indexed with cryptographic SHA-256 integrity digests "
            "to guarantee non-repudiation and chain-of-custody for regulatory examination."
        )
        p_audit_trail.paragraph_format.space_after = Pt(20)

        # Signature Table
        sig_table = doc.add_table(rows=2, cols=2)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sig_c0, sig_c1 = sig_table.rows[0].cells[0], sig_table.rows[0].cells[1]
        sig_c0.width = Inches(3.2)
        sig_c1.width = Inches(3.2)
        
        sig_c0.paragraphs[0].add_run("Lead Cybersecurity Auditor:\n\n_______________________________\nCertified Information Systems Auditor").font.size = Pt(9)
        sig_c1.paragraphs[0].add_run("Authorizing Security Officer:\n\n_______________________________\nChief Information Security Officer").font.size = Pt(9)

        # Save to memory stream
        stream = io.BytesIO()
        doc.save(stream)
        stream.seek(0)
        return stream

